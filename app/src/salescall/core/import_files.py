"""依頼者のファイル（CSV・xlsx）を、行と欄に読む。持ち主 CL-2。DESIGN.md 3-4。

ファイルを読むのは、この部品だけ（発信対象の取込も、これから作る発信禁止の一覧の取込も、ここを通す）。
読むだけで、値の意味（番号か・管理番号か）は判断しない（番号の判断は Q1、行の合否は取込の1か所）。

決まり（推測で読まない。読めないものは、理由をつけて断る）：
- 1行目が見出し。行番号は、表計算ソフトで見える行番号（1行目＝見出し。値の中の改行は1行のまま）。
- CSV の文字の決まり：先頭に BOM があれば UTF-8。無ければ UTF-8・Shift_JIS（cp932）・EUC-JP のどれでも厳密に読み、
  1つだけ読めればそれ、どれで読んでも中身が同じ（英数字だけ）なら UTF-8。
  **2つ以上で読めて中身が違うときは、推測で選ばない**（2026-10-16 の見直し I12：前は EUC-JP を確かめず、EUC-JP の
  ファイルが Shift_JIS として「読めて」しまい、化けた社名のまま通りえた）（`EncodingAmbiguous`。両方の読み方の先頭の数行を返すので、
  画面で見せて契約者が選び、`encoding` を指定して読み直す）。
  測った結果（2026-10-14）：日本語が2文字だけの UTF-8 は約5割が Shift_JIS としても読め、Shift_JIS の2文字は約2%が
  UTF-8 としても読める。⇒「UTF-8 を先に試す」では、短いファイルで名前が化ける（会話で化けた社名を読み上げうる）。
  日本語が100文字を超えると、ほぼ両方では読めなくなるので、ふつうの一覧では選ぶ場面は出ない。
- CSV で、列の数が見出しと合わない行は、その行に印をつける（値の中の「,」が囲まれていないと列がずれ、
  別の列の番号を電話番号として読む恐れがあるため。取込がその行を止める）。
- xlsx：シートが1つならそれ。2つ以上なら、どのシートかを名前で指定しないと断る（黙って1枚目を選ばない）。
  数の型の欄は「数」の印をつける（数は先頭の 0 を持てないので、番号の先頭の 0 が消えた確かな証拠。DESIGN.md 3-4）。
  エラーの値（#N/A など）・結果が保存されていない式・15桁を超える数（表計算ソフトで桁が失われる）は、欄に「問題」をつける。
- 末尾の空の行は読まない（書式だけ残った行）。途中の空の行は、空の行として残す（行番号をずらさない）。
"""

from __future__ import annotations

import codecs
import csv
import hashlib
import io
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException


class FileRejected(Exception):
    """ファイルを読めない（理由は日本語）。何も登録しない。"""


class EncodingAmbiguous(FileRejected):
    """2つ以上の文字の決まり（ENCODINGS）で読めて、中身が違う。samples：読み方ごとの先頭の数行（画面で見せて選んでもらう）。"""

    def __init__(self, samples: dict[str, list[str]]) -> None:
        super().__init__(f"文字の決まり（{' か '.join(_NAMES[e] for e in samples)}）を決められません。正しく読める方を選んでください")
        self.samples = samples


@dataclass(frozen=True)
class Cell:
    text: str | None              # 欄の文字（空・空白だけなら None）
    numeric: bool = False         # xlsx で数の型だった
    problem: str | None = None    # 値を信じてはいけない理由（エラーの値・結果の無い式・桁が失われた数・TRUE/FALSE）
    flag: bool | None = None      # xlsx で TRUE・FALSE の値だった（チェックの列。どの欄で使えるかは取込が決める）


@dataclass(frozen=True)
class Row:
    row_no: int
    cells: tuple[Cell, ...]
    width_ok: bool = True         # CSV で、列の数が見出しと同じ

    def cell(self, index: int) -> Cell:
        return self.cells[index] if index < len(self.cells) else Cell(None)

    @property
    def is_empty(self) -> bool:
        """中身の無い行（チェックの列の外れ＝FALSE も中身ではない。表計算ソフトのチェックの列は、下の空の行まで FALSE が入るので、
        数えると末尾の空の行が「番号の無い行」として並ぶ＝2026-10-18 の見直し）。"""
        return all(c.text is None and (c.problem is None or c.flag is False) for c in self.cells)


@dataclass(frozen=True)
class Sheet:
    file_name: str
    file_kind: str                # csv / xlsx
    file_sha256: str
    encoding: str | None          # CSV の文字の決まり
    sheet_name: str | None        # xlsx のシートの名前
    header: tuple[str | None, ...]
    rows: tuple[Row, ...]


_EMPTY = Cell(None)
_MAX_EXACT = 10 ** 15             # 表計算ソフトの数は15桁まで正確（それより大きい数は、桁が失われている）


ENCODINGS = ("utf-8", "cp932", "euc_jp")
_NAMES = {"utf-8": "UTF-8", "cp932": "Shift_JIS", "euc_jp": "EUC-JP"}   # 画面・知らせに出す名前（ENCODINGS と組）


def read(file_name: str, data: bytes, *, sheet_name: str | None = None, encoding: str | None = None) -> Sheet:
    """ファイルを読む。種類は名前の拡張子で決める（.csv / .xlsx）。読めなければ FileRejected。
    encoding：CSV の文字の決まりを契約者が選んだとき（EncodingAmbiguous のあと）だけ渡す。"""
    lower = file_name.lower()
    digest = hashlib.sha256(data).hexdigest()
    if lower.endswith(".csv"):
        if sheet_name is not None:
            raise FileRejected("CSV にはシートがありません")
        text, encoding = _decode(data, encoding)
        header, rows = _csv_rows(text)
        return Sheet(file_name, "csv", digest, encoding, None, header, rows)
    if lower.endswith(".xlsx"):
        if encoding is not None:
            raise FileRejected("Excel のファイルには、文字の決まりの指定は要りません")
        name, header, rows = _xlsx_rows(data, sheet_name)
        return Sheet(file_name, "xlsx", digest, None, name, header, rows)
    raise FileRejected("読めるのは CSV（.csv）と Excel（.xlsx）だけです（古い .xls は .xlsx で保存し直してください）")


def _text(value: str | None) -> str | None:
    return value if value is not None and value.strip() != "" else None


# ── CSV ──
def _decode(data: bytes, chosen: str | None) -> tuple[str, str]:
    if chosen is not None and chosen not in ENCODINGS:
        raise FileRejected(f"文字の決まりは {' か '.join(_NAMES[e] for e in ENCODINGS)} のどれかです")
    if data.startswith(codecs.BOM_UTF8):
        if chosen not in (None, "utf-8"):
            raise FileRejected("UTF-8 の印（BOM）があるファイルです")
        try:
            text, encoding = data[len(codecs.BOM_UTF8):].decode("utf-8"), "utf-8"
        except UnicodeDecodeError:
            raise FileRejected("UTF-8 の印（BOM）があるのに、UTF-8 として読めません") from None
    else:
        read_as: dict[str, str] = {}
        for enc in ((chosen,) if chosen else ENCODINGS):
            try:
                read_as[enc] = data.decode(enc)
            except UnicodeDecodeError:
                pass
        if not read_as:
            raise FileRejected("選んだ文字の決まりでは読めません" if chosen else
                               f"文字の決まりが {'・'.join(_NAMES[e] for e in ENCODINGS)} のどれでもありません"
                               "（UTF-8 で保存し直してください）")
        if len(set(read_as.values())) > 1:   # 2つ以上の読み方で読めて、中身が違う：推測で選ばない
            raise EncodingAmbiguous({enc: [ln for ln in t.splitlines() if ln.strip()][:5] for enc, t in read_as.items()})
        encoding = next(enc for enc in ENCODINGS if enc in read_as)   # どれで読んでも同じ（英数字だけ）なら、並びの先
        text = read_as[encoding]
    if "\x00" in text:
        raise FileRejected("CSV の中に、文字ではないデータがあります（別の種類のファイルではありませんか）")
    return text, encoding


def _csv_rows(text: str) -> tuple[tuple[str | None, ...], tuple[Row, ...]]:
    try:
        records = list(csv.reader(io.StringIO(text, newline=""), strict=True))
    except csv.Error as e:
        raise FileRejected(f"CSV の書き方が正しくありません（{e}）") from None
    if not records:
        raise FileRejected("ファイルが空です")
    header = tuple(_text(h.strip()) for h in records[0])
    rows = []
    for i, rec in enumerate(records[1:], start=2):
        cells = tuple(Cell(_text(v)) for v in rec)
        row = Row(i, cells, width_ok=len(rec) == len(header))
        if row.is_empty:
            row = Row(i, cells, width_ok=True)   # 空の行（何も無い行）は列の数を問わない
        rows.append(row)
    return header, _trim(rows)


# ── xlsx ──
def _xlsx_rows(data: bytes, sheet_name: str | None) -> tuple[str, tuple[str | None, ...], tuple[Row, ...]]:
    try:
        values = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as e:
        raise FileRejected(f"Excel のファイルとして読めません（{type(e).__name__}）") from None
    try:
        names = [ws.title for ws in values.worksheets if hasattr(ws, "iter_rows")]
        if sheet_name is None:
            if len(names) != 1:
                raise FileRejected(f"シートが {len(names)} 枚あります。どのシートを読むかを選んでください")
            sheet_name = names[0]
        elif sheet_name not in names:
            raise FileRejected("指定のシートがありません")
        ws_v, ws_f = values[sheet_name], formulas[sheet_name]
        ws_v.reset_dimensions()   # ファイルに書かれた範囲を信じない（範囲が誤っているファイルで、行を読み落とさない）
        ws_f.reset_dimensions()
        all_rows = []
        for vr, fr in zip(ws_v.iter_rows(min_row=1, min_col=1), ws_f.iter_rows(min_row=1, min_col=1), strict=True):
            width = max(len(vr), len(fr))
            all_rows.append(tuple(_xlsx_cell(vr[i] if i < len(vr) else None, fr[i] if i < len(fr) else None)
                                  for i in range(width)))
    finally:
        values.close()
        formulas.close()
    if not all_rows:
        raise FileRejected("シートが空です")
    header = tuple(_text(c.text.strip()) if c.text is not None else None for c in all_rows[0])
    rows = [Row(i, cells) for i, cells in enumerate(all_rows[1:], start=2)]
    return sheet_name, header, _trim(rows)


def _xlsx_cell(v, f) -> Cell:
    value = getattr(v, "value", None)
    if getattr(v, "data_type", None) == "e":
        return Cell(None, problem="表計算ソフトのエラーの値です（#N/A など）")
    if value is None:
        if getattr(f, "data_type", None) == "f":
            return Cell(None, problem="式の結果が保存されていません（表計算ソフトで開いて保存し直してください）")
        return _EMPTY
    if isinstance(value, bool):
        return Cell(None, problem="TRUE・FALSE の値です", flag=value)
    if isinstance(value, int | float):
        if abs(value) >= _MAX_EXACT:
            return Cell(None, numeric=True, problem="15桁を超える数で、表計算ソフトで桁が失われています。列を文字にして保存し直してください")
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        return Cell(str(value), numeric=True)
    if isinstance(value, datetime | date | time):
        return Cell(value.isoformat())
    return Cell(_text(str(value)))


def _trim(rows: list[Row]) -> tuple[Row, ...]:
    while rows and rows[-1].is_empty:
        rows.pop()
    return tuple(rows)
