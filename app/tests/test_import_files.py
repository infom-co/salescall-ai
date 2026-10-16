"""ファイルの読み込みの試験（DESIGN.md 3-4）：推測で読まず、読めないものは理由をつけて断る。行番号は表計算ソフトのまま。

⚠️ 中身は架空のもの（番号は局番のあとが 0000 の、実在しない番号）。
"""

from __future__ import annotations

import io
import unittest

import openpyxl

from salescall.core.import_files import EncodingAmbiguous, FileRejected, read


def xlsx(rows, *, sheets=("一覧",)) -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = sheets[0]
    for name in sheets[1:]:
        wb.create_sheet(name)
    ws = wb[sheets[0]]
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(row=r, column=c, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def texts(sheet):
    width = len(sheet.header)
    return [(r.row_no, [r.cell(i).text for i in range(width)]) for r in sheet.rows]


class Csv(unittest.TestCase):
    BODY = "会社名,電話番号\n架空商事,03-0000-0001\n"

    def test_utf8_with_and_without_bom_and_shift_jis(self):
        long_body = self.BODY + "".join(f"架空の株式会社その{i},03-0000-{i:04d}\n" for i in range(2, 30))
        for data, enc in ((b"\xef\xbb\xbf" + self.BODY.encode("utf-8"), "utf-8"),
                          (long_body.encode("utf-8"), "utf-8"),
                          (long_body.encode("cp932"), "cp932")):
            with self.subTest(enc=enc):
                s = read("一覧.csv", data)
                self.assertEqual((s.encoding, s.header), (enc, ("会社名", "電話番号")))
                self.assertEqual(texts(s)[0], (2, ["架空商事", "03-0000-0001"]))

    def test_ambiguous_encoding_is_not_guessed(self):
        """短い日本語は、UTF-8 とも Shift_JIS とも読めることがある。推測で選ばず、両方を見せて契約者が選ぶ。"""
        data = "名前,番号\nあい,03-0000-0001\n".encode("utf-8")
        with self.assertRaises(EncodingAmbiguous) as cm:
            read("一覧.csv", data)
        self.assertEqual(set(cm.exception.samples), {"utf-8", "cp932"})
        self.assertIn("あい,03-0000-0001", cm.exception.samples["utf-8"])
        self.assertEqual(read("一覧.csv", data, encoding="utf-8").header, ("名前", "番号"))
        self.assertNotEqual(read("一覧.csv", data, encoding="cp932").header, ("名前", "番号"))

    def test_ascii_only_is_utf8(self):
        self.assertEqual(read("a.csv", b"name,phone\nx,03-0000-0001\n").encoding, "utf-8")

    def test_unreadable_files_are_rejected(self):
        for name, data in (("a.csv", b""), ("a.csv", b"\xff\xfe\x00a"), ("a.csv", b"a,b\n\x00,1\n"),
                           ("a.xls", b"x"), ("a.txt", b"a"), ("a.xlsx", b"not a zip"),
                           ("a.csv", b'a,b\n"x,1\n')):
            with self.subTest(name=name, data=data):
                with self.assertRaises(FileRejected):
                    read(name, data)

    def test_row_numbers_follow_the_spreadsheet(self):
        data = 'a,b\nx,1\n\n"改行\nのある値",2\ny,3\n\n\n'.encode("utf-8")
        s = read("a.csv", data)
        self.assertEqual([r.row_no for r in s.rows], [2, 3, 4, 5])   # 値の中の改行は1行のまま。末尾の空の行は読まない
        self.assertTrue(s.rows[1].is_empty)

    def test_width_mismatch_is_marked(self):
        s = read("a.csv", "会社名,電話番号\n架空商事,東京支店,03-0000-0001\n架空工業\nx,03-0000-0002\n".encode("utf-8"))
        self.assertEqual([r.width_ok for r in s.rows], [False, False, True])

    def test_cells_keep_their_text(self):
        s = read("a.csv", b"a,b\n  ,'0300000001\n")
        self.assertEqual(texts(s), [(2, [None, "'0300000001"])])   # 空白だけは空。値は手を加えない（番号の判断は Q1）


class Xlsx(unittest.TestCase):
    def test_numbers_are_marked_numeric(self):
        s = read("a.xlsx", xlsx([["電話番号", "管理番号"], [300000001, "A-1"], ["03-0000-0002", 1001], [300000003.0, None]]))
        self.assertEqual(texts(s), [(2, ["300000001", "A-1"]), (3, ["03-0000-0002", "1001"]), (4, ["300000003", None])])
        self.assertEqual([[r.cell(i).numeric for i in range(2)] for r in s.rows], [[True, False], [False, True], [True, False]])

    def test_untrustworthy_cells_carry_a_problem(self):
        s = read("a.xlsx", xlsx([["電話番号"], ["#N/A"], ["=1+1"], [1234567890123456], [True]]))
        self.assertEqual([r.cells[0].text for r in s.rows], [None, None, None, None])
        self.assertTrue(all(r.cells[0].problem for r in s.rows))
        self.assertFalse(any(r.is_empty for r in s.rows))

    def test_sheet_must_be_chosen_when_there_are_several(self):
        data = xlsx([["電話番号"], ["03-0000-0001"]], sheets=("一覧", "メモ"))
        with self.assertRaises(FileRejected):
            read("a.xlsx", data)
        self.assertEqual(read("a.xlsx", data, sheet_name="一覧").sheet_name, "一覧")
        with self.assertRaises(FileRejected):
            read("a.xlsx", data, sheet_name="無いシート")

    def test_row_numbers_and_empty_rows(self):
        s = read("a.xlsx", xlsx([["電話番号"], ["03-0000-0001"], [None], ["03-0000-0002"], [None], [None]]))
        self.assertEqual([(r.row_no, r.is_empty) for r in s.rows], [(2, False), (3, True), (4, False)])


if __name__ == "__main__":
    unittest.main()
