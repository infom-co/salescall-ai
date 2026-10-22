"""業務設定の取込（引き継ぎの (10)。core/business_import.py）の試験。架空の業務。"""

from __future__ import annotations

import io
import unittest

import openpyxl

from salescall.core import business_import as bi
from salescall.core import business_settings as bs
from salescall.core import settings
from tests.support import OPERATOR, World

REASONS = ("担当ではない", "まだ決まっていない", "いまは分からない", "理由を言わなかった")


def book(*, sheets=None, drop=(), rename=None) -> bytes:
    """標準のひな形どおりの、架空の業務設定のファイル。sheets で中身を差し替え、drop でシートを外す。"""
    content = {
        "基本": [("項目", "値"), ("業務の目的", "予約の希望日時を伺う"), ("お客様の区分", "法人"), ("自社名", "株式会社テスト物流"),
               ("自社名の読み", "かぶしきがいしゃてすとぶつりゅう"), ("担当者名", "山田"), ("担当者名の読み", "やまだ"),
               ("用件の一文", "ご予約の確認でお電話いたしました。"), ("目的のお客様", "誰でも"), ("目的のお客様の役職", None),
               ("忙しいと言われたときの一言", None), ("声の型", "S2")],
        "質問": [("質問ID", "質問の文", "聞くこと", "型", "必須", "選択肢", "断りの答え"),
               ("Q1", "「テスト便」のご予約はなさいますか。", "ご予約", "選択肢", "はい", "する／しない", "しない"),
               ("Q2", "ご希望の日時をお聞かせください。", "ご希望の日時", "日時", "はい", None, None)],
        "分岐": [("質問ID", "回答条件", "次の質問ID"), ("Q1", "しない", "終わり"), ("Q1", "その他", "Q2")],
        "わからない・答えたくない": [("質問ID", "答え", "理由", "扱い", "結論の名前")]
        + [("全部", "わからない", r, "日を改める", None) for r in REASONS]
        + [("全部", "答えたくない", r, "結論として終える", "回答不可") for r in ("決まりで言えない", "答えたくない")],
        "かけ直し": [("結果", "かけ直すか", "回数", "間隔")]
        + [(k, "かける", "3", "2時間おき") for k in ("会話に至らなかった", "一部だけ取れた", "不在")]
        + [("かけ直しの時刻が決まっていない", "かける", "2", "1日後")],
        "言ってはいけないこと": [("名前", "言ってはいけないこと", "正しい案内"), ("F1", "初月は無料です", "初月は半額です")],
        "名前": [("表記", "読み"), ("テスト便", "てすとびん")],
        "想定問答": [("問いの名前", "答え"), ("番号の出どころ", "以前ご予約の際にいただいたお電話番号にお掛けしております。")],
    }
    content.update(sheets or {})
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in content.items():
        if name in drop:
            continue
        ws = wb.create_sheet((rename or {}).get(name, name))
        for r in rows:
            ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def book_rows(sheet: str) -> list[tuple]:
    """標準のひな形の、そのシートの行（試験で書き換えるため）。"""
    wb = openpyxl.load_workbook(io.BytesIO(book()))
    return [tuple(c.value for c in r) for r in wb[sheet].iter_rows()]


def finder(*items: dict):
    """名前を洗い出す処理の代わり（CL-1 の names.decide の戻り値の形を返す。聞かれた台本を残す）。"""
    def find(text: str) -> list[dict]:
        find.asked.append(text)
        return [dict(x) for x in items]
    find.asked = []
    return find


class Import(unittest.TestCase):
    def setUp(self):
        self.w = World()
        self.s = self.w.setup()

    def tx(self, work):
        return self.w.a.in_client(OPERATOR, self.s["tenant"], self.s["client"], work)

    def preview(self, data, **kw):
        return self.tx(lambda tx: bi.preview(tx, self.s["business"], "業務設定.xlsx", data, **kw))

    def start_verification(self):
        """検証モードの業務（AI の読みの案を契約者が確かめる。切り替えは契約者の承認と同じ1つ）。"""
        self.tx(lambda tx: settings.set_value(tx, "approval.human", True, scope="business", status="decided",
                                              rationale="試験：検証モード", review_when="試験",
                                              tenant_id=self.s["tenant"], client_id=self.s["client"],
                                              business_id=self.s["business"]))

    def test_preview_and_register(self):
        """入口が持ってきた読みは、契約者が確かめずに使う（管理者も契約者も介在しない。確かめるのは検証モードの AI の案だけ）。"""
        data = book()
        p = self.preview(data)
        self.assertEqual((p.problems, p.names_to_confirm), ([], []))
        v = self.tx(lambda tx: bi.register(tx, self.s["business"], "業務設定.xlsx", data, expected_digest=p.digest))
        got = self.tx(lambda tx: bs.load(tx, v))
        self.assertEqual((got.next_of("Q1", "しない"), got.rules.decline["Q1"], got.rules.retry["callback_unspecified_time"].interval),
                         ("END", frozenset({"しない"}), ("days", 1)))
        self.assertEqual(got.rules.handling[("Q2", "refused", "policy")].conclusion, "回答不可")   # 「全部」の行が効く
        self.assertIn(bs.Name("テスト便", "てすとびん", "script", "given"), got.names)
        self.assertEqual(got.answers, (("番号の出どころ", "以前ご予約の際にいただいたお電話番号にお掛けしております。"),))   # 想定問答

    def test_register_refuses_a_changed_preview(self):
        p = self.preview(book())
        other = book(sheets={"基本": [("項目", "値"), ("業務の目的", "目的を書き換えた"), ("お客様の区分", "法人"), ("自社名", "株式会社テスト物流"),
                                     ("自社名の読み", "かぶしきがいしゃてすとぶつりゅう"), ("担当者名", "山田"),
                                     ("担当者名の読み", "やまだ"), ("目的のお客様", "誰でも"), ("声の型", "S2")]})
        with self.assertRaises(bi.PreviewChanged):
            self.tx(lambda tx: bi.register(tx, self.s["business"], "業務設定.xlsx", other, expected_digest=p.digest))

    def test_problems_name_the_sheet_and_row(self):
        bad = book(sheets={"質問": [("質問ID", "質問の文", "聞くこと", "型", "必須", "選択肢", "断りの答え"),
                                   ("Q1", "ご予約はなさいますか。", None, "色", "はい", None, None)]}, drop=("分岐",))
        p = self.preview(bad)
        self.assertTrue(any("シート「質問」の 2 行目の型" in x for x in p.problems), p.problems)

    def test_missing_required_sheet_and_mapping(self):
        p = self.preview(book(drop=("かけ直し",)))
        self.assertTrue(any("かけ直し" in x for x in p.problems), p.problems)
        renamed = book(rename={"質問": "質問一覧"})
        p = self.preview(renamed, mapping={"sheets": {"questions": "質問一覧"}})
        self.assertEqual(p.problems, [])                                  # 依頼者の形を対応付けて読む

    def test_names_found_in_the_script_are_added(self):
        """台本の名前は、名前を洗い出す処理（CL-1 の names.decide）が見つける。入口の名前に無いものも、名前の一覧に入る。"""
        found = finder({"surface": "テスト便", "reading": "てすとびん", "source": "モデル2つ", "candidates": []},
                       {"surface": "中島倉庫", "reading": "", "source": "", "candidates": ["なかじまそうこ", "なかしまそうこ"],
                        "likely": "なかじまそうこ"})
        p = self.preview(book(sheets={"名前": [("表記", "読み")]}), find_names=found)
        self.assertEqual(p.problems, [])
        self.assertIn("テスト便", found.asked[0])                      # 台本（質問の文など）をまとめて1回
        got = {n["text"]: n for n in p.raw["names"]}
        self.assertEqual((got["テスト便"]["reading"], got["テスト便"]["reading_by"]), ("てすとびん", "ai"))
        self.assertEqual((got["中島倉庫"]["reading"], got["中島倉庫"]["likely"]), (None, "なかじまそうこ"))

    def test_an_undecided_reading_does_not_stop_the_import(self):
        """読みが決まらない名前でも止めない。考えられる読みとまず言ってみる読みを持って登録し、会話の中で確かめる（2026-10-16
        ユーザー：「なかじまとなかしまなどは人が見ても判断できません」）。検証モードでも、契約者の確かめの対象にしない。"""
        self.start_verification()
        data = book(sheets={"名前": [("表記", "読み"), ("テスト便", None)]})
        found = finder({"surface": "テスト便", "reading": "", "source": "", "candidates": ["てすとびん", "てすとべん"],
                        "likely": "てすとびん"})
        p = self.preview(data, find_names=found)
        self.assertEqual((p.problems, p.names_to_confirm), ([], []))
        v = self.tx(lambda tx: bi.register(tx, self.s["business"], "業務設定.xlsx", data, expected_digest=p.digest, ai=p.ai))
        self.assertIn(bs.Name("テスト便", None, "script", None, ("てすとびん", "てすとべん"), "てすとびん"),
                      self.tx(lambda tx: bs.load(tx, v)).names)
        p = self.preview(data)                                            # 名前を洗い出す処理が答えなくても止めない
        self.assertEqual(p.problems, [])

    def test_during_verification_a_person_confirms_the_ai_reading(self):
        """検証モードでは、AI が決めた読みを契約者が確かめてから（直してもよい）。2026-10-16 ユーザー：「検証中は
        AIが認識したものを私が確認します」。"""
        self.start_verification()
        data = book(sheets={"名前": [("表記", "読み"), ("テスト便", None)]})
        found = finder({"surface": "テスト便", "reading": "てすとべん", "source": "モデル2つ", "candidates": []})   # AI（誤り）
        p = self.preview(data, find_names=found)
        self.assertEqual(p.names_to_confirm, [{"text": "テスト便", "reading": "てすとべん", "source": "script"}])
        self.assertTrue(any("AI が決めた読みです" in x for x in p.problems), p.problems)
        fixed = [("テスト便", "てすとびん")]                                # 契約者が「てすとびん」に直して確かめた
        p = self.preview(data, find_names=found, confirmed_names=fixed)
        self.assertEqual((p.problems, p.names_to_confirm), ([], []))
        v = self.tx(lambda tx: bi.register(tx, self.s["business"], "業務設定.xlsx", data, expected_digest=p.digest,
                                           confirmed_names=fixed, ai=p.ai))
        self.assertIn(bs.Name("テスト便", "てすとびん", "script", "person"), self.tx(lambda tx: bs.load(tx, v)).names)

    def test_reviewing_a_preview_again_keeps_who_decided_the_reading(self):
        """プレビューの形を、もう一度確かめても（画面で別の所を直して出し直すなど）、AI の読みは AI の読みのまま。
        検証モードの確かめは抜けない（2026-10-16 の見直し I1：前は入口の読みに化けて、確かめなしで登録できた）。"""
        self.start_verification()
        data = book(sheets={"名前": [("表記", "読み"), ("テスト便", None)]})
        found = finder({"surface": "テスト便", "reading": "てすとべん", "source": "モデル2つ", "candidates": []})
        p1 = self.preview(data, find_names=found)
        p2 = self.tx(lambda tx: bi.review(tx, self.s["business"], p1.raw))          # AI の答えを渡さずに出し直す
        self.assertEqual(p2.names_to_confirm, [{"text": "テスト便", "reading": "てすとべん", "source": "script"}])
        self.assertTrue(any("AI が決めた読みです" in x for x in p2.problems), p2.problems)
        by = {n["text"]: n["reading_by"] for n in p2.raw["names"]}
        self.assertEqual(by["テスト便"], "ai")

    def test_the_entrance_reading_of_the_caller_is_not_replaced_by_the_ai(self):
        """名乗りの自社名に入口の読みがあれば、名前の一覧の同じ表記も、その読み（AI の読みで上書きしない。2026-10-16 の見直し I3）。
        名乗りと名前の一覧で入口の読みが食い違えば、問題に出す。"""
        from tests.test_business_settings import raw
        r = raw()
        r["caller"]["company"] = {"text": "東海", "reading": "とうみ"}
        r["names"] = [n for n in r["names"] if n["text"] != "株式会社テスト物流"] + [
            {"text": "東海", "reading": None, "source": "script"}]
        ai = {"names": {"東海": {"reading": "とうかい", "reading_by": "ai", "candidates": [], "likely": None}}}
        p = self.tx(lambda tx: bi.review(tx, self.s["business"], r, ai=ai))
        got = {n["text"]: (n["reading"], n["reading_by"]) for n in p.raw["names"]}
        self.assertEqual((got["東海"], p.raw["caller"]["company"]["reading"]), (("とうみ", "given"), "とうみ"))
        r["names"][-1]["reading"] = "とうかい"
        p = self.tx(lambda tx: bi.review(tx, self.s["business"], r, ai=ai))
        self.assertTrue(any("名乗りと名前の一覧で違います" in x for x in p.problems), p.problems)

    def test_after_verification_the_ai_reading_is_used_without_a_person(self):
        """フルオート（検証モードでない業務）では、AI が決めた読みをそのまま使う（管理者も契約者も介在しない）。台本の読みがなは
        入口の読み。どれで決めたかは残る。"""
        data = book(sheets={"名前": [("表記", "読み"), ("テスト便", None)]})
        found = finder({"surface": "テスト便", "reading": "てすとびん", "source": "モデル2つ", "candidates": []},
                       {"surface": "東海林", "reading": "しょうじ", "source": "台本の読みがな", "candidates": []})
        p = self.preview(data, find_names=found)
        self.assertEqual((p.problems, p.names_to_confirm), ([], []))
        v = self.tx(lambda tx: bi.register(tx, self.s["business"], "業務設定.xlsx", data, expected_digest=p.digest,
                                           ai=p.ai))
        names = self.tx(lambda tx: bs.load(tx, v)).names
        self.assertIn(bs.Name("テスト便", "てすとびん", "script", "ai"), names)
        self.assertIn(bs.Name("東海林", "しょうじ", "script", "given"), names)

    def test_the_callers_name_needs_its_reading_from_the_entrance(self):
        """担当者名の読みは、入口で必ず入れる（2026-10-16 ユーザー決定＝案C）。AI が読みを決めても使わない。"""
        rows = [r for r in book_rows("基本") if r[0] != "担当者名の読み"]
        found = finder({"surface": "山田", "reading": "やまだ", "source": "モデル2つ", "candidates": []})
        p = self.preview(book(sheets={"基本": rows}), find_names=found)
        self.assertTrue(any("担当者名の読みは必ず入れます" in x for x in p.problems), p.problems)

    def test_reviewing_again_does_not_make_the_callers_ai_reading_an_entrance_reading(self):
        """名乗りの読みは、名前の一覧の読みの写し。もう一度確かめたとき、写した AI の読みを入口の読みにしない（2026-10-17 の見直し：
        前は担当者名の読みが無くても、2回目の確かめで問題が消え、AI の読みが入口の読みとして登録できた）。"""
        rows = [r for r in book_rows("基本") if r[0] != "担当者名の読み"]
        found = finder({"surface": "山田", "reading": "やまだ", "source": "モデル2つ", "candidates": []})
        p1 = self.preview(book(sheets={"基本": rows}), find_names=found)
        p2 = self.tx(lambda tx: bi.review(tx, self.s["business"], p1.raw, ai=p1.ai))
        self.assertTrue(any("担当者名の読みは必ず入れます" in x for x in p2.problems), p2.problems)
        self.assertEqual({n["text"]: n["reading_by"] for n in p2.raw["names"]}["山田"], "ai")
        self.start_verification()                                       # 検証モード：自社名の AI の読みの確かめも抜けない
        from tests.test_business_settings import raw
        r = raw()
        r["caller"]["company"]["reading"] = None
        r["names"] = [dict(n, reading=None) if n["text"] == "株式会社テスト物流" else n for n in r["names"]]
        ai = {"names": {"株式会社テスト物流": {"reading": "かぶしきがいしゃてすとぶつりゅう", "reading_by": "ai",
                                          "candidates": [], "likely": None}}}
        q1 = self.tx(lambda tx: bi.review(tx, self.s["business"], r, ai=ai))
        q2 = self.tx(lambda tx: bi.review(tx, self.s["business"], q1.raw))
        self.assertEqual([n["text"] for n in q2.names_to_confirm], ["株式会社テスト物流"])

    def test_malformed_input_is_a_problem_not_a_crash(self):
        """画面・抜き出しの入口から崩れた形が来ても、落ちずに問題として出す（2026-10-17 の見直し）。"""
        from tests.test_business_settings import raw
        for decline in (None, 5, "しない"):
            with self.subTest(decline=decline):
                r = raw()
                r["rules"]["decline"] = {"Q1": decline}
                p = self.tx(lambda tx: bi.review(tx, self.s["business"], r))
                self.assertTrue(any("断りの答え" in x for x in p.problems), p.problems)
                self.assertFalse(any("聞かずに終わる" in x for x in p.problems), p.problems)   # うその理由を足さない
        r = raw()
        r["names"].append({"text": 123, "reading": "いち", "source": "script"})
        p = self.tx(lambda tx: bi.review(tx, self.s["business"], r))
        self.assertTrue(any("表記" in x for x in p.problems), p.problems)
        retry = [("結果", "かけ直すか", "回数", "間隔")] + [
            (k, "かける", "²", "2時間おき") for k in ("会話に至らなかった", "一部だけ取れた", "不在")
        ] + [("かけ直しの時刻が決まっていない", "かける", "2", "1日後")]
        p = self.preview(book(sheets={"かけ直し": retry}))
        self.assertTrue(any("回数（1以上の数）" in x for x in p.problems), p.problems)

    def test_the_old_not_transferred_retry_row_is_skipped_with_a_note(self):
        """前のひな形の「取り次いでもらえなかった」の行は、システムが決める（日を改めて1回）ので、問題にせず読み飛ばして知らせる
        （⛔ 2026-10-22 の見直し13：前は「使えない結果」として登録を止めた）。"""
        from salescall.core import business_rules
        retry = book_rows("かけ直し") + [("取り次いでもらえなかった", "かける", "5", "2時間おき")]
        p = self.preview(book(sheets={"かけ直し": retry}))
        self.assertEqual(p.problems, [])
        self.assertTrue(any("取り次いでもらえなかった" in x and "システムが決めるので読み飛ばしました（日を改めて1回）" in x
                            for x in p.notes), p.notes)
        self.assertNotIn("not_reached_not_transferred", p.raw["rules"]["retry"])
        self.assertEqual(business_rules.SYSTEM_POLICIES["not_reached_not_transferred"],
                         business_rules.Policy("call", 1, ("days", 1)))        # 知らせの「日を改めて1回」と同じ
        retry = book_rows("かけ直し") + [("取り次いでもらえない", "かける", "1", "1日後")]
        p = self.preview(book(sheets={"かけ直し": retry}))
        self.assertTrue(any("取り次いでもらえない" in x for x in p.problems), p.problems)   # 打ち違いは、これまでどおり問題

    def test_a_missing_absent_row_takes_the_first_value(self):
        """「不在」（つながったが目的のお客様が不在で、戻る時刻が取れなかった）の行が無ければ、初めの値（かける・3回）で入れ、
        知らせに出す（2026-10-22 ユーザー：「つながったけど不在だった場合は、依頼者ごとに選択できるようにしておき、デフォルトは
        3回ぐらい」）。書いてあれば、その値。初めの値の無い結果（会話に至らなかった＝電話がつながらなかった など）は、これまで
        どおり書かないと問題になる。行があって形が違えば、初めの値で埋めずに問題に出す。"""
        from salescall.core import business_rules
        rows = [r for r in book_rows("かけ直し") if r[0] != "不在"]
        p = self.preview(book(sheets={"かけ直し": rows}))
        self.assertEqual(p.problems, [])
        self.assertEqual(p.raw["rules"]["retry"]["not_reached_absent"],
                         business_rules.policy_raw(business_rules.DEFAULT_POLICIES["not_reached_absent"]))
        self.assertEqual(business_rules.DEFAULT_POLICIES["not_reached_absent"].max_times, 3)
        self.assertTrue(any("「不在」の行が無いので、初めの値（かける・3回・3時間おき）で入れました" in x for x in p.notes), p.notes)
        p = self.preview(book())                                   # 書いてあればその値（3回・2時間おき）で、知らせも出さない
        self.assertEqual(p.raw["rules"]["retry"]["not_reached_absent"],
                         {"action": "call", "max_times": 3, "interval": {"hours": 2}})
        self.assertFalse(any("初めの値" in x for x in p.notes), p.notes)
        rows = [r for r in book_rows("かけ直し") if r[0] != "会話に至らなかった"]
        p = self.preview(book(sheets={"かけ直し": rows}))
        self.assertTrue(any("再発信の方針" in x for x in p.problems), p.problems)
        rows = [r if r[0] != "不在" else ("不在", "かける", "x", "2時間おき") for r in book_rows("かけ直し")]
        p = self.preview(book(sheets={"かけ直し": rows}))
        self.assertTrue(any("回数（1以上の数）" in x for x in p.problems), p.problems)
        self.assertFalse(any("初めの値" in x for x in p.notes), p.notes)

    def test_problems_found_outside_the_check_are_in_the_digest(self):
        """確かめ（check）の外で見つかった問題（重なった行など。中身には残らない）は指紋に入る：プレビューから問題を消しても
        登録できない（2026-10-17 の見直し）。"""
        rows = book_rows("名前") + [("テスト便", "てすとびん")]                 # 同じ名前の行が2つ
        p = self.preview(book(sheets={"名前": rows}))
        self.assertTrue(p.problems)
        p.problems = []
        with self.assertRaises(bi.PreviewChanged):
            self.tx(lambda tx: bi.register_reviewed(tx, self.s["business"], p, expected_digest=p.digest))

    def test_a_version_with_an_unconfirmed_ai_reading_is_refused_in_verification_mode(self):
        """プレビューを通さずに版を作っても、同じ確かめ（業務設定の1か所）が止める。"""
        from tests.test_business_settings import raw
        r = raw()
        r["names"][2]["reading_by"] = "ai"
        self.assertIsNotNone(self.tx(lambda tx: bs.save_version(tx, self.s["business"], r)))   # フルオート
        self.start_verification()
        with self.assertRaises(bs.SettingsInvalid):
            self.tx(lambda tx: bs.save_version(tx, self.s["business"], r))

    def test_the_ai_is_asked_once_and_registration_uses_the_previewed_answer(self):
        """AI は同じ問いに毎回同じ答えを返すとは限らない。プレビューで1回だけ聞き、登録はその答えを使う（聞き直すと、契約者が
        見たものと違う中身になり、登録できなくなる。2026-10-16 に見つけた穴）。"""
        data = book(sheets={"名前": [("表記", "読み"), ("テスト便", None)]})
        answers = iter(["てすとびん", "てすとべん"])                 # 2回目は違う答え
        asked = []

        def ai(text):
            asked.append(text)
            return [{"surface": "テスト便", "reading": next(answers), "source": "モデル2つ", "candidates": []}]
        p = self.preview(data, find_names=ai)
        v = self.tx(lambda tx: bi.register(tx, self.s["business"], "業務設定.xlsx", data, expected_digest=p.digest, ai=p.ai))
        self.assertEqual(len(asked), 1)                                   # 登録では聞かない
        self.assertIn(bs.Name("テスト便", "てすとびん", "script", "ai"), self.tx(lambda tx: bs.load(tx, v)).names)

    def test_a_reading_must_be_kana(self):
        """AI が漢字やローマ字を返したら、読みとして使わない（読みが決まらない側に倒す）。入口の読みが使えなければ問題に出す。
        読みはひらがなで持つ（カタカナはひらがなに直す）。"""
        data = book(sheets={"名前": [("表記", "読み"), ("テスト便", None)]})
        p = self.preview(data, find_names=finder({"surface": "テスト便", "reading": "test bin", "source": "モデル2つ",
                                                  "candidates": []}))
        self.assertEqual(p.problems, [])
        self.assertEqual({n["text"]: n["reading"] for n in p.raw["names"]}["テスト便"], None)
        bad = book(sheets={"名前": [("表記", "読み"), ("テスト便", "手素都便")]})
        p = self.preview(bad)
        self.assertTrue(any("ひらがなかカタカナ" in x for x in p.problems), p.problems)
        kata = book(sheets={"名前": [("表記", "読み"), ("テスト便", "テストビン")]})
        self.assertEqual({n["text"]: n["reading"] for n in self.preview(kata).raw["names"]}["テスト便"], "てすとびん")

    def test_missing_optional_sheet_is_noted_and_a_mapped_one_is_a_problem(self):
        """無くてもよいシートが無いときは、プレビューに知らせる（黙って読み飛ばさない）。対応付けで名前を指定したシートが
        無ければ問題（2026-10-16 の見直し I5）。"""
        p = self.preview(book(rename={"言ってはいけないこと": "禁止事項"}))
        self.assertEqual(p.problems, [])
        self.assertIn("シート「言ってはいけないこと」がありません（無いものとして読みました。シートの名前を確かめてください）", p.notes)
        p = self.preview(book(drop=("名前",)), mapping={"sheets": {"names": "名前の一覧"}})
        self.assertTrue(any("シート「名前の一覧」" in x for x in p.problems), p.problems)

    def test_a_choice_named_other_can_be_branched_and_the_star_is_the_catch_all(self):
        """選択肢に「その他」があれば、分岐の「その他」はその選択肢。受け皿は「＊」（2026-10-16 の見直し I6）。"""
        q = [("質問ID", "質問の文", "聞くこと", "型", "必須", "選択肢", "断りの答え"),
             ("Q1", "どちらになさいますか。", "ご希望", "選択肢", "はい", "A／B／その他", None),
             ("Q2", "詳しくお聞かせください。", "詳しく", "自由な意見", "いいえ", None, None),
             ("Q3", "ご希望の日時をお聞かせください。", "ご希望の日時", "日時", "はい", None, None)]
        b = [("質問ID", "回答条件", "次の質問ID"), ("Q1", "その他", "Q2"), ("Q1", "＊", "Q3")]
        p = self.preview(book(sheets={"質問": q, "分岐": b}))
        nxt = {x["id"]: x["next"] for x in p.raw["questions"]}
        self.assertEqual(nxt["Q1"], {"その他": "Q2", bs.ANY: "Q3"})

    def test_duplicate_rows_are_problems_not_silently_overwritten(self):
        """同じ結果のかけ直し・同じ言ってはいけないこと・同じ質問×答え×理由の行が2つあれば、問題に出す（2026-10-16 の見直し I7）。"""
        retry = book_rows("かけ直し") + [("不在", "終える", None, None)]
        forbidden = book_rows("言ってはいけないこと") + [("F1", "必ず安くなります", None)]
        handling = book_rows("わからない・答えたくない") + [("全部", "わからない", REASONS[0], "結論として終える", "不明")]
        p = self.preview(book(sheets={"かけ直し": retry, "言ってはいけないこと": forbidden, "わからない・答えたくない": handling}))
        for text in ("結果「不在」の行が2つあります", "名前「F1」が2つあります", "の行が2つあります（"):
            self.assertTrue(any(text in x for x in p.problems), (text, p.problems))

    def test_problems_point_to_the_sheet_and_row(self):
        """名前・質問の問題は、ファイルのシートと行で出す（前は自社名・担当者名の分ずれて「名前 3 番目」と出た。I10）。"""
        q = book_rows("質問")
        q.insert(2, (None, None, None, None, None, None, None))                # 空の行（数えない）
        q[3] = (None,) + q[3][1:]                                               # 4行目＝2つ目の質問の番号が無い
        p = self.preview(book(sheets={"質問": q}))
        self.assertTrue(any(x.startswith("シート「質問」の 4 行目の番号") for x in p.problems), p.problems)
        names = {"名前": ["シート「基本」の自社名", "シート「基本」の担当者名", "シート「名前」の 2 行目"]}
        self.assertEqual(bi._placed(["名前 3 番目の表記 が空です", "名前 4 番目の読み"], names),
                         ["シート「名前」の 2 行目の表記 が空です", "名前 4 番目の読み"])   # ファイルに無い名前（AI が足した）はそのまま

    def test_answer_problems_point_to_the_sheet_and_row(self):
        """想定問答の問題も、シートと行で出す。想定問答のシートは無くてもよい（CL-1 の連絡帳 4-26）。"""
        rows = book_rows("想定問答") + [("番号:出どころ", "以前お電話いただいた番号です。")]
        p = self.preview(book(sheets={"想定問答": rows}))
        self.assertTrue(any(x.startswith("シート「想定問答」の 3 行目の問いの名前") for x in p.problems), p.problems)
        self.assertEqual(self.preview(book(drop=("想定問答",))).problems, [])

    def test_registration_recomputes_the_digest_and_includes_the_business(self):
        """登録は、プレビューの指紋を信じず中身から出し直す。別の業務の指紋では登録できない（2026-10-16 の見直し I8）。"""
        p = self.preview(book())
        p.raw["purpose"] = "書き換えた目的"                                  # プレビューを書き換えても
        with self.assertRaises(bi.PreviewChanged):
            self.tx(lambda tx: bi.register_reviewed(tx, self.s["business"], p, expected_digest=p.digest))
        other, _ = self.w.business_with_version(self.s["tenant"], self.s["client"])
        q = self.preview(book())
        with self.assertRaises(bi.PreviewChanged):
            self.tx(lambda tx: bi.register_reviewed(tx, other, q, expected_digest=q.digest))

    def test_units_of_number_questions_are_decided_without_a_person(self):
        """数の質問の単位は、単位を決める処理（CL-1 の1か所）の答えを入れる。管理者も契約者も確かめない。決まらなければ空。"""
        rows = [("質問ID", "質問の文", "聞くこと", "型", "必須", "選択肢", "断りの答え"),
                ("Q1", "「テスト便」のご予約はなさいますか。", "ご予約", "選択肢", "はい", "する／しない", "しない"),
                ("Q2", "ご希望の日時をお聞かせください。", "ご希望の日時", "日時", "はい", None, None),
                ("Q3", "ご利用は何名様でしょうか。", "人数", "数", "はい", None, None),
                ("Q4", "拠点の数を教えていただけますか。", "拠点の数", "数", "いいえ", None, None)]
        data = book(sheets={"質問": rows, "分岐": [("質問ID", "回答条件", "次の質問ID"), ("Q1", "しない", "終わり"),
                                                  ("Q1", "その他", "Q2")]})
        asked = []

        def decide(texts):
            asked.extend(texts)
            return {"ご利用は何名様でしょうか。": "名", "拠点の数を教えていただけますか。": ""}   # 2つ目は決まらない
        p = self.preview(data, decide_units=decide)
        self.assertEqual(p.problems, [])
        self.assertEqual(asked, ["ご利用は何名様でしょうか。", "拠点の数を教えていただけますか。"])   # 数の質問だけ
        v = self.tx(lambda tx: bi.register(tx, self.s["business"], "業務設定.xlsx", data, expected_digest=p.digest,
                                           ai=p.ai))
        got = {q.id: q.unit for q in self.tx(lambda tx: bs.load(tx, v)).questions}
        self.assertEqual(got, {"Q1": None, "Q2": None, "Q3": "名", "Q4": None})

    def test_what_to_ask_about_is_decided_without_a_person(self):
        """質問の「聞くこと」は、入口に無ければ CL-1 の決め方（topics.decide）の答えを入れる。決まらなければ空。"""
        rows = [("質問ID", "質問の文", "聞くこと", "型", "必須", "選択肢", "断りの答え"),
                ("Q1", "「テスト便」のご予約はなさいますか。", "ご予約", "選択肢", "はい", "する／しない", "しない"),
                ("Q2", "ご希望の日時をお聞かせください。", None, "日時", "はい", None, None)]
        data = book(sheets={"質問": rows})
        asked = []

        def decide(texts):
            asked.extend(texts)
            return {"ご希望の日時をお聞かせください。": "ご希望の日時"}
        p = self.preview(data, decide_topics=decide)
        self.assertEqual((p.problems, asked), ([], ["ご希望の日時をお聞かせください。"]))   # 入口に無い質問だけ
        v = self.tx(lambda tx: bi.register(tx, self.s["business"], "業務設定.xlsx", data, expected_digest=p.digest, ai=p.ai))
        got = {q.id: q.ask_about for q in self.tx(lambda tx: bs.load(tx, v)).questions}
        self.assertEqual(got, {"Q1": "ご予約", "Q2": "ご希望の日時"})

    def test_meanings_and_replies_are_decided_without_a_person(self):
        """はい／いいえ・選択肢の質問の意味の言葉と曖昧な言葉の意味は、取込のときに AI（CL-1 の決め方）に1回ずつだけ聞いて入れる
        （管理者も契約者も確かめない）。使えない組（宣言に無い選択肢・決まった一覧に無い言葉）は入れない。答えが無い・形が違うときは
        空のまま登録できる（会話は止まらない。CL-1 の連絡帳 4-17）。登録では聞き直さない。"""
        asked = []

        def meanings(pairs):
            asked.append(("meanings", pairs))
            return [{"する": "ご予約なさるということ", "しない": "ご予約はなさらないということ", "わからない": "x"} for _ in pairs]

        def replies(pairs):
            asked.append(("replies", pairs))
            return [{"結構": "しない", "たぶん": "する"} for _ in pairs]

        data = book()
        p = self.preview(data, decide_meanings=meanings, decide_replies=replies)
        self.assertEqual(p.problems, [])
        self.assertEqual(asked, [("meanings", [("「テスト便」のご予約はなさいますか。", ("する", "しない"))]),
                                 ("replies", [("「テスト便」のご予約はなさいますか。", ("する", "しない"))])])
        v = self.tx(lambda tx: bi.register(tx, self.s["business"], "業務設定.xlsx", data, expected_digest=p.digest, ai=p.ai))
        q1 = self.tx(lambda tx: bs.load(tx, v)).questions[0]
        self.assertEqual((q1.meanings, q1.replies),
                         ({"する": "ご予約なさるということ", "しない": "ご予約はなさらないということ"}, {"結構": "しない"}))
        self.assertEqual(len(asked), 2)                                   # 登録では聞き直さない
        p = self.preview(data, decide_meanings=lambda pairs: [], decide_replies=None)   # 形の違う答え・聞けない
        self.assertEqual((p.problems, p.raw["questions"][0]["meanings"], p.raw["questions"][0]["replies"]), ([], None, None))

    def test_other_entrances_go_the_same_way(self):
        """依頼者の画面から抜き出した中身（同じ形）も、同じ確かめと登録の道を通る。"""
        from tests.test_business_settings import raw
        extracted = raw()
        extracted["names"].append({"text": "テスト便", "reading": None, "source": "script"})
        p = self.tx(lambda tx: bi.review(tx, self.s["business"], extracted))
        self.assertEqual(p.problems, [])                                   # 読みが決まらなくても止めない（会話で確かめる）
        v = self.tx(lambda tx: bi.register_reviewed(tx, self.s["business"], p, expected_digest=p.digest))
        self.assertEqual(self.tx(lambda tx: bs.load(tx, v)).company.text, "株式会社テスト物流")

    def test_nothing_is_written_on_problems(self):
        bad = book(drop=("かけ直し",))
        p = self.preview(bad)
        with self.assertRaises(bs.SettingsInvalid):
            self.tx(lambda tx: bi.register(tx, self.s["business"], "業務設定.xlsx", bad, expected_digest=p.digest))


if __name__ == "__main__":
    unittest.main()
